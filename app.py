import streamlit as st
from google import genai
from google.genai import types
import html
import io
import json
import re
from docx import Document
import PyPDF2
from fpdf import FPDF

# --- Constantes Globais ---
CHAR_REPLACEMENTS = {
    'č': 'c', 'ć': 'c', 'ž': 'z', 'š': 's', 'đ': 'd',
    'Č': 'C', 'Ć': 'C', 'Ž': 'Z', 'Š': 'S', 'Đ': 'D',
    'ł': 'l', 'ą': 'a', 'ę': 'e', 'ó': 'o', 'ś': 's', 'ź': 'z', 'ż': 'z',
    'Ł': 'L', 'Ą': 'A', 'Ę': 'E', 'Ó': 'O', 'Ś': 'S', 'Ź': 'Z', 'Ż': 'Z',
    '\u2013': '-', '\u2014': '-',
    '\u201c': '"', '\u201d': '"',
    '\u2018': "'", '\u2019': "'",
    '\u2022': '-', '•': '-'
}

URL_PATTERN = re.compile(r'((?:https?://)?(?:www\.)?(?:github\.com|linkedin\.com)[^\s|]*)', re.IGNORECASE)

MODEL_MAPPING = {
    "Gemini 3.5 Flash (Recomendado)": "gemini-3.5-flash",
    "Gemini 3 Flash": "gemini-3.0-flash",
    "Gemini 2.5 Flash": "gemini-2.5-flash",
    "Gemma 4 31B": "gemma-4-31b-it",
    "Gemma 4 26B": "gemma-4-26b-it"
}

SECTION_CAPS = {
    "requirements": 40,
    "experience": 30,
    "terms": 20,
    "education": 10,
}

ATS_SCHEMA = {
    "type": "object",
    "properties": {
        "overall_score": {"type": "integer"},
        "keywords_present": {"type": "array", "items": {"type": "string"}},
        "keywords_missing": {"type": "array", "items": {"type": "string"}},
        "section_scores": {
            "type": "object",
            "properties": {
                "requirements": {"type": "integer"},
                "experience": {"type": "integer"},
                "terms": {"type": "integer"},
                "education": {"type": "integer"},
            },
            "required": ["requirements", "experience", "terms", "education"],
        },
        "diagnostic": {"type": "string"},
        "next_steps": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "impact": {"type": "string", "enum": ["high", "medium", "low"]},
                    "title": {"type": "string"},
                    "bullets": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["impact", "title", "bullets"],
            },
        },
    },
    "required": [
        "overall_score",
        "keywords_present",
        "keywords_missing",
        "section_scores",
        "diagnostic",
        "next_steps",
    ],
}

# --- Leitura de Ficheiros ---
def read_uploaded(uploaded_file):
    if uploaded_file is None:
        return None
    bytes_data = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()
    try:
        if filename.endswith(('.txt', '.md')):
            return bytes_data.decode('utf-8')
        elif filename.endswith('.pdf'):
            reader = PyPDF2.PdfReader(io.BytesIO(bytes_data))
            return "\n".join(page.extract_text() for page in reader.pages if page.extract_text())
        elif filename.endswith('.docx'):
            document = Document(io.BytesIO(bytes_data))
            return "\n".join(p.text for p in document.paragraphs)
    except Exception as e:
        st.error(f"Erro ao ler '{uploaded_file.name}': {e}")
    return None

# --- PDF Harvard ---
class HarvardPDF(FPDF):
    def header(self): pass
    def footer(self): pass

def sanitize_text(text):
    for u, a in CHAR_REPLACEMENTS.items(): text = text.replace(u, a)
    return text

def draw_section_line(pdf, y=None):
    if y is None: y = pdf.get_y()
    pdf.set_draw_color(0, 0, 0)
    pdf.set_line_width(0.4)
    pdf.line(15, y, 195, y)

def render_url_line(pdf, line, line_height):
    segments = URL_PATTERN.split(line)
    if '|' in line:
        total_w = sum(pdf.get_string_width(seg) for seg in segments if seg)
        start_x = 15 + (180 - total_w) / 2
        pdf.set_x(start_x)
    for seg in segments:
        if not seg: continue
        if URL_PATTERN.fullmatch(seg):
            url = seg if seg.startswith('http') else f'https://{seg}'
            pdf.set_text_color(17, 85, 204)
            pdf.cell(pdf.get_string_width(seg), line_height, seg, link=url)
            pdf.set_text_color(0, 0, 0)
        else:
            pdf.cell(pdf.get_string_width(seg), line_height, seg)
    pdf.ln(line_height)

def convert_markdown_to_harvard_pdf(md_text):
    md_text = sanitize_text(md_text)
    pdf = HarvardPDF()
    pdf.add_page()
    pdf.set_margins(15, 12, 15)
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.set_font("Times", size=10)
    lines = md_text.split('\n')
    header_block_ended = False
    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(1.2)
            continue
        line = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', line)
        if line.startswith('# '):
            name = line[2:].replace('**', '').replace('*', '')
            pdf.set_font("Times", "B", size=16)
            pdf.cell(180, 7, name, align='C')
            pdf.ln(7)
            header_block_ended = False
        elif line.startswith('## '):
            if not header_block_ended:
                pdf.ln(1)
                draw_section_line(pdf)
                pdf.ln(2)
                header_block_ended = True
            section = line[3:].replace('**', '').replace('*', '').upper()
            pdf.ln(2)
            pdf.set_font("Times", "B", size=10)
            pdf.cell(180, 4, section)
            pdf.ln(5)
            draw_section_line(pdf)
            pdf.ln(2)
        elif line.startswith('### '):
            sub = line[4:].replace('**', '').replace('*', '')
            pdf.set_font("Times", "B", size=9)
            pdf.cell(180, 4, sub)
            pdf.ln(4.5)
            header_block_ended = True
        elif line.startswith(('* ', '- ')):
            bullet_text = line[2:].replace('**', '').replace('*', '').replace('`', '')
            pdf.set_font("Times", size=9)
            pdf.set_x(20)
            pdf.cell(4, 3.8, "-", ln=False)
            pdf.multi_cell(171, 3.8, bullet_text)
            pdf.set_x(15)
            header_block_ended = True
        else:
            clean = line.replace('**', '').replace('*', '').replace('`', '')
            pdf.set_font("Times", size=9)
            if URL_PATTERN.search(clean):
                render_url_line(pdf, clean, 3.5)
            elif '|' in clean:
                pdf.cell(180, 3.5, clean, align='C')
                pdf.ln(3.5)
            else:
                pdf.multi_cell(180, 3.8, clean)
                pdf.set_x(15)
    return bytes(pdf.output())

# --- CSS Corrigido ---
def inject_custom_css():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:ital,wght@0,400;0,600;1,400&family=Inter:wght@300;400;500;600&family=IBM+Plex+Mono:wght@400;500;600&display=swap');
    
    :root {
        --bg-primary: #09090b;       
        --bg-card: #18181b;          
        --bg-card-hover: #27272a;    
        --border: #27272a;           
        --accent: #818cf8;           
        --accent-light: #a5b4fc;
        --success: #34d399;          
        --warning: #fbbf24;          
        --danger: #f87171;           
        --text-primary: #fafafa;     
        --text-secondary: #a1a1aa;   
        --text-muted: #52525b;       
    }
    
    html, body, [class*="css-"] { font-family: 'Inter', sans-serif !important; color: var(--text-primary) !important; }
    .stApp { background: var(--bg-primary) !important; }
    h1, h2, h3, h4 { font-family: 'Playfair Display', serif !important; color: var(--text-primary) !important; font-weight: 400 !important; letter-spacing: -0.02em; }
    
    /* Native Buttons */
    .stButton > button { 
        background: var(--bg-card) !important; color: var(--text-primary) !important; 
        border-radius: 12px !important; border: 1px solid var(--border) !important;
        padding: 14px 20px !important; transition: all 0.2s ease !important; 
        box-shadow: 0 1px 2px rgba(0,0,0,0.2) !important; width: 100% !important;
    }
    .stButton > button:hover { 
        border-color: var(--accent) !important; background: var(--bg-card-hover) !important;
        transform: translateY(-1px); box-shadow: 0 4px 12px rgba(0,0,0,0.4) !important; 
    }
    button[data-testid="baseButton-primary"] {
        background: var(--accent) !important; color: #fff !important; font-weight: 500 !important; border: none !important;
        box-shadow: 0 4px 14px rgba(129, 140, 248, 0.3) !important;
    }
    button[data-testid="baseButton-primary"]:hover { background: #6366f1 !important; transform: translateY(-2px); }
    
    /* Componentes de UI */
    .detail-card { background: var(--bg-card); border: 1px solid var(--border); border-radius: 16px; padding: 24px; box-shadow: 0 4px 20px rgba(0,0,0,0.2); height: 100%; }
    .section-badge { font-family: 'IBM Plex Mono', monospace; font-size: 0.7rem; font-weight: 600; color: var(--accent); text-transform: uppercase; letter-spacing: 1.5px; margin-bottom: 12px; }
    .keyword-tag-present { display: inline-block; background: rgba(52, 211, 153, 0.1); color: var(--success); border: 1px solid rgba(52, 211, 153, 0.2); border-radius: 6px; padding: 4px 10px; font-size: 0.8rem; margin: 3px; font-family: 'IBM Plex Mono', monospace; }
    .keyword-tag-missing { display: inline-block; background: rgba(248, 113, 113, 0.1); color: var(--danger); border: 1px solid rgba(248, 113, 113, 0.2); border-radius: 6px; padding: 4px 10px; font-size: 0.8rem; margin: 3px; font-family: 'IBM Plex Mono', monospace; }
    .next-step-block { background: var(--bg-card); border: 1px solid var(--border); border-radius: 14px; padding: 22px 26px; margin-bottom: 14px; }
    .impact-badge { font-family: 'IBM Plex Mono', monospace; font-size: 0.65rem; padding: 4px 10px; border-radius: 6px; font-weight: 600; margin-left: 10px; text-transform: uppercase; }
    .impact-high { background: rgba(52, 211, 153, 0.15); color: var(--success); border: 1px solid rgba(52, 211, 153, 0.3); }
    .impact-medium { background: rgba(251, 191, 36, 0.15); color: var(--warning); border: 1px solid rgba(251, 191, 36, 0.3); }
    
    .progress-track { width: 100%; height: 6px; background: var(--bg-primary); border-radius: 3px; overflow: hidden; margin-top: 12px; border: 1px solid var(--border); }
    .progress-fill { height: 100%; border-radius: 3px; transition: width 0.8s ease-out; }
    
    /* Wizard Nav */
    .wizard-progress { width: 100%; height: 3px; background: var(--border); border-radius: 3px; margin-bottom: 0; overflow: hidden; }
    .wizard-progress-bar { height: 100%; background: var(--accent); border-radius: 3px; transition: width 0.4s ease; }
    .wizard-step-indicator { font-family: 'IBM Plex Mono', monospace; font-size: 0.8rem; color: var(--text-muted); text-align: center; margin-top: 8px; letter-spacing: 1px; }
    
    /* Overrides */
    .stTextArea textarea { background: var(--bg-primary) !important; border: 1px solid var(--border) !important; border-radius: 8px !important; color: var(--text-primary) !important; }
    .stTextArea textarea:focus { border-color: var(--accent) !important; box-shadow: none !important; }
    .stFileUploader>div { background: var(--bg-primary) !important; border: 1px dashed var(--border) !important; border-radius: 12px !important; }
    
    /* Esconder headers por default */
    .stDeployButton, header[data-testid="stHeader"] { display: none !important; }
    </style>
    """, unsafe_allow_html=True)

# --- Helpers ---
def esc(value) -> str:
    return html.escape(str(value), quote=True)


def validate_ats_data(raw: dict) -> dict:
    required = {
        "overall_score",
        "keywords_present",
        "keywords_missing",
        "section_scores",
        "diagnostic",
        "next_steps",
    }
    missing = required - raw.keys()
    if missing:
        raise ValueError(f"Resposta incompleta da IA: faltam {sorted(missing)}")

    sections = {}
    for key, cap in SECTION_CAPS.items():
        sections[key] = max(0, min(int(raw["section_scores"].get(key, 0)), cap))

    next_steps = []
    for step in raw.get("next_steps", [])[:3]:
        if not isinstance(step, dict):
            continue
        impact = step.get("impact", "medium")
        if impact not in ("high", "medium", "low"):
            impact = "medium"
        bullets = [str(b) for b in step.get("bullets", []) if b][:5]
        next_steps.append({
            "impact": impact,
            "title": str(step.get("title", "")),
            "bullets": bullets,
        })

    return {
        "overall_score": max(0, min(int(raw["overall_score"]), 100)),
        "keywords_present": [str(k) for k in raw.get("keywords_present", [])][:20],
        "keywords_missing": [str(k) for k in raw.get("keywords_missing", [])][:20],
        "section_scores": sections,
        "diagnostic": str(raw.get("diagnostic", "")),
        "next_steps": next_steps,
    }


def get_hex_color(pct):
    if pct >= 70: return "#34d399"
    if pct >= 40: return "#fbbf24"
    return "#f87171"

def gauge_svg(score, max_score=100):
    pct = min(score / max_score, 1.0)
    color = get_hex_color(pct * 100)
    r, cx, cy = 80, 100, 100
    circumference = 2 * 3.14159 * r
    dash = circumference * pct
    gap = circumference - dash
    return (
        f'<svg viewBox="0 0 200 200" xmlns="http://www.w3.org/2000/svg" style="width:220px;height:220px;">'
        f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="#27272a" stroke-width="12"/>'
        f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="{color}" stroke-width="12" stroke-linecap="round" stroke-dasharray="{dash:.1f} {gap:.1f}" transform="rotate(-90 {cx} {cy})"/>'
        f'<text x="{cx}" y="{cy - 5}" text-anchor="middle" font-family="Playfair Display,serif" font-size="44" fill="#fafafa">{score}</text>'
        f'<text x="{cx}" y="{cy + 20}" text-anchor="middle" font-family="IBM Plex Mono,monospace" font-size="10" fill="#a1a1aa" letter-spacing="1">COM A VAGA</text>'
        f'</svg>'
    )

def subscore_circle(value, max_val, label):
    pct = value / max_val if max_val > 0 else 0
    color = get_hex_color(pct * 100)
    r = 28
    circ = 2 * 3.14159 * r
    dash = circ * pct
    gap = circ - dash
    return (
        f'<div style="display:flex;flex-direction:column;align-items:center;gap:8px;padding:10px;background:rgba(24,24,27,0.5);border-radius:12px;border:1px solid #27272a;">'
        f'<svg viewBox="0 0 70 70" style="width:60px;height:60px;">'
        f'<circle cx="35" cy="35" r="{r}" fill="none" stroke="#27272a" stroke-width="5"/>'
        f'<circle cx="35" cy="35" r="{r}" fill="none" stroke="{color}" stroke-width="5" stroke-linecap="round" stroke-dasharray="{dash:.1f} {gap:.1f}" transform="rotate(-90 35 35)"/>'
        f'<text x="35" y="40" text-anchor="middle" font-family="IBM Plex Mono,monospace" font-size="14" font-weight="600" fill="#fafafa">{value}</text>'
        f'</svg>'
        f'<div style="font-family:\'IBM Plex Mono\',monospace;font-size:0.6rem;letter-spacing:1px;text-transform:uppercase;color:#a1a1aa;text-align:center;">{label}</div>'
        f'</div>'
    )

def progress_bar_html(value, max_val, color=None):
    pct = min(value / max_val, 1.0) * 100 if max_val > 0 else 0
    c = color or get_hex_color(pct)
    return (
        f'<div class="progress-track">'
        f'<div class="progress-fill" style="width:{pct:.1f}%;background:{c};"></div>'
        f'</div>'
        f'<div style="font-family:\'IBM Plex Mono\',monospace;font-size:0.75rem;color:{c};letter-spacing:1px;margin-top:8px;">{int(pct)}% DO CRITÉRIO</div>'
    )

# --- App Setup ---
st.set_page_config(layout="wide", page_title="ATS AI Resume Optimizer")
inject_custom_css()

# Session State Init (Começa logo no Step 0)
for key, default in [
    ("wizard_step", 0), ("q1_answer", None), ("q2_answer", None), ("q3_answer", None),
    ("jd_text_pasted", ""), ("suggestions", []), 
    ("ats_data", None), ("final_markdown", ""), ("cv_content", None), ("job_content", None)
]:
    if key not in st.session_state:
        st.session_state[key] = default

def get_client(): return genai.Client(api_key=google_api_key)

# Sidebar
with st.sidebar:
    st.header("Configuração")
    google_api_key = st.text_input("Google AI Studio API Key", type="password")
    selected_model_label = st.selectbox("Modelo de IA:", options=list(MODEL_MAPPING.keys()))
    selected_model = MODEL_MAPPING[selected_model_label]
    if st.button("Reiniciar App"):
        for key in ["wizard_step", "q1_answer", "q2_answer", "q3_answer", "jd_text_pasted",
                    "suggestions", "ats_data", "final_markdown", "cv_content", "job_content"]:
            st.session_state.pop(key, None)
        st.rerun()

# --- WIZARD QUESTIONS (steps 0-2) ---
WIZARD_QUESTIONS = [
    {
        "tag": "PASSO 1", "title": "Estás a ser chamado para entrevistas?",
        "subtitle": "A tua taxa de resposta ajuda a calibrar a agressividade da optimização.",
        "var": "q1_answer",
        "options": [
            {"title": "Quase nunca", "desc": "Raramente recebo respostas ou convites.", "val": "Quase nunca"},
            {"title": "Às vezes", "desc": "Recebo algumas respostas, mas poderia ser melhor.", "val": "Às vezes"},
            {"title": "Com frequência", "desc": "Tenho entrevistas regulares, quero optimizar mais.", "val": "Com frequência"},
        ],
    },
    {
        "tag": "PASSO 2", "title": "Qual é o teu objetivo principal?",
        "subtitle": "Queremos personalizar as sugestões ao teu momento de carreira.",
        "var": "q2_answer",
        "options": [
            {"title": "Primeiro emprego", "desc": "Estou a começar e preciso de destacar potencial.", "val": "Primeiro emprego"},
            {"title": "Mudar de empresa", "desc": "Quero dar o próximo passo na trajetória.", "val": "Mudar empresa"},
            {"title": "Melhorar o CV atual", "desc": "Tenho experiência, quero um CV mais forte.", "val": "Melhorar CV atual"},
        ],
    },
    {
        "tag": "PASSO 3", "title": "Que tipo de vaga procuras?",
        "subtitle": "O nível de senioridade altera os pesos da análise ATS.",
        "var": "q3_answer",
        "options": [
            {"title": "Estágio / Junior", "desc": "Exigem menos experiência e mais formação.", "val": "Estágio / Junior"},
            {"title": "Pleno / Sénior", "desc": "Focam-se em resultados e stack técnica.", "val": "Pleno / Sénior"},
            {"title": "Sem vaga específica", "desc": "Quero um CV generalista forte.", "val": "Sem vaga específica"},
        ],
    },
]

if 0 <= st.session_state.wizard_step < 3:
    q_idx = st.session_state.wizard_step
    q = WIZARD_QUESTIONS[q_idx]
    progress_pct = int(((q_idx + 1) / 4) * 100)

    st.markdown(f"""
    <div style="padding-top:8px;padding-bottom:24px;">
      <div class="wizard-progress"><div class="wizard-progress-bar" style="width:{progress_pct}%;"></div></div>
      <div class="wizard-step-indicator">{q_idx + 1} / 4</div>
    </div>
    <div style="max-width:620px;margin:0 auto;padding-top:16px;text-align:center;">
      <div class="section-badge" style="justify-content:center; display:flex;">{q["tag"]}</div>
      <div style="font-size:2.2rem;margin-bottom:12px; font-family:'Playfair Display',serif;">{q["title"]}</div>
      <div style="font-size:1rem;color:var(--text-secondary);margin-bottom:36px;">{q["subtitle"]}</div>
    </div>
    """, unsafe_allow_html=True)

    _, center_col, _ = st.columns([1, 2, 1])
    with center_col:
        for opt in q["options"]:
            if st.button(f"**{opt['title']}** — {opt['desc']}", key=f"wiz_{q_idx}_{opt['val']}", use_container_width=True):
                st.session_state[q["var"]] = opt["val"]
                st.session_state.wizard_step += 1
                st.rerun()
                
        if q_idx > 0:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("← Voltar ao passo anterior", key="back"):
                st.session_state.wizard_step -= 1
                st.rerun()
    st.stop()

# --- STEP 3: Upload ---
if st.session_state.wizard_step == 3:
    st.markdown("""
    <div style="padding-top:8px;padding-bottom:24px;">
      <div class="wizard-progress"><div class="wizard-progress-bar" style="width:100%;"></div></div>
      <div class="wizard-step-indicator">4 / 4</div>
    </div>
    <div style="max-width:800px;margin:0 auto;text-align:center;">
      <div class="section-badge" style="justify-content:center; display:flex;">ÚLTIMO PASSO</div>
      <div style="font-size:2.2rem;margin-bottom:12px; font-family:'Playfair Display',serif;">Carrega os teus dados</div>
      <div style="font-size:1rem;color:var(--text-secondary);margin-bottom:28px;">Faz upload do CV e da descrição da vaga.</div>
    </div>
    """, unsafe_allow_html=True)

    c1, c2 = st.columns([1, 1], gap="large")
    with c1:
        st.markdown("<div class='section-badge'>TEU CURRÍCULO</div>", unsafe_allow_html=True)
        cv_file = st.file_uploader("Upload do CV", type=["pdf", "docx", "txt", "md"], label_visibility="collapsed")
        if cv_file: 
            st.session_state.cv_content = read_uploaded(cv_file)
        
        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("← Voltar"):
            st.session_state.wizard_step -= 1
            st.rerun()

    with c2:
        st.markdown("<div class='section-badge'>DESCRIÇÃO DA VAGA</div>", unsafe_allow_html=True)
        # TABS PARA RECUPERAR O DRAG AND DROP
        tab1, tab2 = st.tabs(["📋 Colar Texto", "📁 Upload / Drag & Drop"])
        
        with tab1:
            st.session_state.jd_text_pasted = st.text_area(
                "Cola a descrição da vaga", value=st.session_state.jd_text_pasted, height=200, label_visibility="collapsed",
                placeholder="Cola aqui os requisitos, tecnologias e responsabilidades..."
            )
            if st.session_state.jd_text_pasted.strip():
                st.session_state.job_content = st.session_state.jd_text_pasted.strip()
                
        with tab2:
            jd_file = st.file_uploader("Upload da Vaga", type=["pdf", "docx", "txt", "md"], label_visibility="collapsed", key="jd_uploader")
            if jd_file:
                st.session_state.job_content = read_uploaded(jd_file)

    ready = bool(st.session_state.cv_content) and bool(st.session_state.job_content)
    
    st.markdown("<br>", unsafe_allow_html=True)
    _, col_btn, _ = st.columns([1, 1, 1])
    with col_btn:
        if st.button("Iniciar Análise ATS", type="primary", use_container_width=True, disabled=not ready):
            if not google_api_key:
                st.error("⚠️ Insere a tua Google AI Studio API Key na barra lateral para prosseguir.")
            else:
                st.session_state.wizard_step = 4
                st.rerun()
    st.stop()


# --- STEP 4: ATS Call ---
if st.session_state.wizard_step == 4:
    if st.session_state.get("ats_data"):
        st.session_state.wizard_step = 5
        st.rerun()

    if not google_api_key:
        st.error("⚠️ Insere a tua Google AI Studio API Key na barra lateral para prosseguir.")
        st.stop()

    with st.spinner("A cruzar o teu CV com os requisitos da vaga..."):
        try:
            client = get_client()
            ats_prompt = f"""You are an expert ATS parser. Analyze this resume vs job description.

Rules:
- keywords_present must only include terms verifiable in the resume text.
- section_scores caps: requirements≤40, experience≤30, terms≤20, education≤10.
- overall_score should approximate the sum of section_scores (0-100).
- next_steps: max 3 items, actionable, no fabrication.

Resume:
{st.session_state.cv_content}

Job description:
{st.session_state.job_content}"""

            response = client.models.generate_content(
                model=selected_model,
                contents=ats_prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=ATS_SCHEMA,
                    temperature=0.1,
                ),
            )
            st.session_state.ats_data = validate_ats_data(json.loads(response.text))
            st.session_state.wizard_step = 5
            st.rerun()
        except Exception as e:
            st.error(f"Erro na API: {e}")
            if st.button("Tentar Novamente"):
                st.rerun()
            st.stop()

# --- STEP 5: Results Dashboard ---
if st.session_state.wizard_step == 5 and st.session_state.ats_data:
    d = st.session_state.ats_data
    overall = int(d.get("overall_score", 0))
    sections = d.get("section_scores", {})
    req_s, exp_s, terms_s, edu_s = sections.get("requirements", 0), sections.get("experience", 0), sections.get("terms", 0), sections.get("education", 0)

    top_left, top_right = st.columns([1, 1], gap="large")
    with top_left:
        st.markdown(gauge_svg(overall, 100), unsafe_allow_html=True)
        c1, c2, c3, c4 = st.columns(4)
        with c1: st.markdown(subscore_circle(req_s, 40, "REQUISITOS"), unsafe_allow_html=True)
        with c2: st.markdown(subscore_circle(exp_s, 30, "EXPERIÊNCIA"), unsafe_allow_html=True)
        with c3: st.markdown(subscore_circle(terms_s, 20, "TERMOS"), unsafe_allow_html=True)
        with c4: st.markdown(subscore_circle(edu_s, 10, "FORMAÇÃO"), unsafe_allow_html=True)

    with top_right:
        ad_cls = "impact-high" if overall >= 70 else ("impact-medium" if overall >= 40 else "")
        ad_lbl = "ALTA ADERÊNCIA" if overall >= 70 else ("MÉDIA" if overall >= 40 else "BAIXA")
        st.markdown(f"""
        <div style="padding-top:10px;">
          <div class="section-badge">DIAGNÓSTICO DA IA</div>
          <div style="font-size:2.2rem; font-family:'Playfair Display',serif; margin-bottom:12px; line-height:1.1;">Como o teu CV é lido</div>
          <div style="margin-bottom:18px;"><span class="impact-badge {ad_cls}">{ad_lbl}</span></div>
          <div style="background:var(--bg-card); border:1px solid var(--border); border-radius:12px; padding:18px; color:var(--text-secondary); line-height:1.5;">
            💡 {esc(d.get("diagnostic", ""))}
          </div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<hr style='border-color:var(--border); margin:40px 0;'>", unsafe_allow_html=True)
    
    col1, col2 = st.columns(2, gap="medium")
    with col1:
        st.markdown(f"""
        <div class="detail-card">
          <div class="section-badge">REQUISITOS (40pts)</div>
          <div style="font-size:1.4rem; font-family:'Playfair Display',serif; margin-bottom:10px;">Cobertura do Core</div>
          {progress_bar_html(req_s, 40)}
        </div>
        """, unsafe_allow_html=True)
        
        ptags = "".join(
            f'<span class="keyword-tag-present">{esc(kw)}</span>'
            for kw in d.get("keywords_present", [])[:8]
        )
        mtags = "".join(
            f'<span class="keyword-tag-missing">{esc(kw)}</span>'
            for kw in d.get("keywords_missing", [])[:8]
        )
        st.markdown(f"""
        <div class="detail-card" style="margin-top:16px;">
          <div class="section-badge">TERMOS DA VAGA (20pts)</div>
          <div style="margin-bottom:12px;">{progress_bar_html(terms_s, 20)}</div>
          <div style="margin-bottom:4px; font-size:0.7rem; color:var(--success);">✓ ENCONTRADOS</div>
          <div style="margin-bottom:12px;">{ptags or '<span style="color:var(--text-muted)">Nenhum</span>'}</div>
          <div style="margin-bottom:4px; font-size:0.7rem; color:var(--danger);">✗ EM FALTA</div>
          <div>{mtags or '<span style="color:var(--text-muted)">Nenhum</span>'}</div>
        </div>
        """, unsafe_allow_html=True)

    with col2:
        st.markdown(f"""
        <div class="detail-card" style="margin-bottom:16px;">
          <div class="section-badge">EXPERIÊNCIA (30pts)</div>
          <div style="font-size:1.4rem; font-family:'Playfair Display',serif; margin-bottom:10px;">Maturidade Profissional</div>
          {progress_bar_html(exp_s, 30)}
        </div>
        <div class="detail-card">
          <div class="section-badge">FORMAÇÃO (10pts)</div>
          <div style="font-size:1.4rem; font-family:'Playfair Display',serif; margin-bottom:10px;">Base Académica</div>
          {progress_bar_html(edu_s, 10)}
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<div class='section-badge' style='margin-top:40px;'>PLANO DE ACÇÃO</div>", unsafe_allow_html=True)
    for i, step in enumerate(d.get("next_steps", [])[:3]):
        imp = step.get("impact", "medium")
        cls = "impact-high" if imp == "high" else ("impact-medium" if imp == "medium" else "")
        lbl = "ALTO IMPACTO" if imp == "high" else "MÉDIO IMPACTO"
        bullets = "".join(
            f"<li style='color:var(--text-secondary); margin-bottom:4px;'>{esc(b)}</li>"
            for b in step.get("bullets", [])
        )
        st.markdown(f"""
        <div class="next-step-block">
          <div style="display:flex; align-items:center; margin-bottom:10px;">
            <span style="color:var(--accent); font-weight:bold; margin-right:8px;">0{i+1}</span>
            <span class="impact-badge {cls}">{lbl}</span>
          </div>
          <div style="font-size:1.2rem; font-family:'Playfair Display',serif; margin-bottom:8px;">{esc(step.get('title', ''))}</div>
          <ul style="margin:0; padding-left:16px;">{bullets}</ul>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    _, btn_col, _ = st.columns([1, 1, 1])
    with btn_col:
        if st.button("Gerar CV Optimizado", type="primary", use_container_width=True):
            st.session_state.wizard_step = 6
            st.rerun()

# --- STEP 6+: Optimização do CV ---
if st.session_state.wizard_step == 6:
    if st.session_state.get("final_markdown"):
        st.session_state.wizard_step = 7
        st.rerun()

    if not google_api_key:
        st.error("⚠️ Insere a tua Google AI Studio API Key na barra lateral para prosseguir.")
        st.stop()

    with st.spinner("A reescrever o teu CV para passar no ATS..."):
        try:
            client = get_client()
            sys_prompt = f"You are an expert HR. Rewrite the resume completely to match the job description. Tailor it for someone whose goal is: {st.session_state.q2_answer or 'career growth'}. Output ONLY clean Markdown starting with '# Name'. No conversational text."
            user_prompt = f"Context: Interview response rate: {st.session_state.q1_answer or 'N/A'}, Goal: {st.session_state.q2_answer or 'N/A'}, Target level: {st.session_state.q3_answer or 'N/A'}.\nOriginal Resume: {st.session_state.cv_content}\nJob: {st.session_state.job_content}"
            response = client.models.generate_content(
                model=selected_model, contents=user_prompt,
                config=types.GenerateContentConfig(system_instruction=sys_prompt, temperature=0.3, max_output_tokens=4000)
            )
            st.session_state.final_markdown = response.text
            st.session_state.wizard_step = 7
            st.rerun()
        except Exception as e:
            st.error(f"Erro na geração: {e}")
            if st.button("Tentar Novamente"):
                st.rerun()
            st.stop()

if st.session_state.wizard_step == 7:
    st.markdown("<div class='section-badge'>CV FINALIZADO</div>", unsafe_allow_html=True)
    st.markdown("<h2 style='margin-bottom: 24px;'>Download do Documento</h2>", unsafe_allow_html=True)
    
    pdf_data = convert_markdown_to_harvard_pdf(st.session_state.final_markdown)
    
    d1, d2, d3 = st.columns([1, 1, 1])
    with d1:
        st.download_button("📥 Descarregar PDF (Harvard)", data=pdf_data, file_name="Optimized_CV.pdf", mime="application/pdf", type="primary", use_container_width=True)
    with d2:
        st.download_button("📥 Descarregar Markdown", data=st.session_state.final_markdown.encode('utf-8'), file_name="CV.md", mime="text/markdown", use_container_width=True)
    with d3:
        if st.button("🔄 Analisar Novo CV", use_container_width=True):
            for key in ["wizard_step", "q1_answer", "q2_answer", "q3_answer", "jd_text_pasted",
                        "suggestions", "ats_data", "final_markdown", "cv_content", "job_content"]:
                st.session_state.pop(key, None)
            st.rerun()
            
    with st.expander("Pré-visualizar CV Gerado", expanded=True):
        st.markdown(st.session_state.final_markdown)